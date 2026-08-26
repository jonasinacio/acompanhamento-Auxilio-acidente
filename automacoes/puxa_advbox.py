#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PUXA-ADVBOX · a boca (fase 2)
=============================

Puxa do AdvBox e ACRESCENTA nas "mães" o que ainda não existe — para as
planilhas se preencherem sozinhas, em vez de à mão. Regra de ouro do ARAUTO:
**casa pelo nº do processo e só acrescenta o que não existe; NUNCA sobrescreve
célula preenchida nem toca em carimbo.**

Contrato real do AdvBox (dos scripts do Jonas no Drive):
  base https://app.advbox.com.br/api/v1 · Authorization: Bearer <token>
  • PERÍCIAS  (PRONTO): GET /settings -> tarefas cujo nome contém "PERICIA";
    GET /posts?task_id=<id> -> tarefas; data/hora/local saem do texto livre
    `notes` (mesmo parser do buscar-advbox-pericias.js).
  • EMENDAS e CASOS (mapeamento PENDENTE): o de-para do AdvBox pra essas mães
    ainda não foi definido — hoje só rodam em --mock. Quando você me disser (ou
    eu ler) de qual tarefa/status vêm, eu ligo igual à perícia.

Uso:
    python3 puxa_advbox.py --dry                       # mostra o que faria (padrão)
    python3 puxa_advbox.py --write --destino pericia    # grava só as perícias
    python3 puxa_advbox.py --dry --mock                 # exemplos, sem token
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
import pj_comum as pj  # noqa: E402

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Rode:  pip3 install openpyxl", file=sys.stderr)
    sys.exit(2)


ADVBOX_ENDPOINT = os.environ.get("ADVBOX_URL", "https://app.advbox.com.br/api/v1").rstrip("/")
ADVBOX_TOKEN    = os.environ.get("ADVBOX_TOKEN", "")
# O AdvBox (WAF) bloqueia requisições sem cara de navegador — mesmo UA dos
# scripts do Jonas (buscar-advbox-*.js), senão devolve 403.
ADVBOX_UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
             "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")


def _get_json(path: str, params: dict | None = None):
    url = f"{ADVBOX_ENDPOINT}/{path.lstrip('/')}"
    if params:
        url += "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={
        "Authorization": f"Bearer {ADVBOX_TOKEN}", "Accept": "application/json",
        "User-Agent": ADVBOX_UA})
    with urllib.request.urlopen(req, timeout=45) as r:
        return json.loads(r.read().decode("utf-8", "replace"))


# ----------------------------------------------------------------------
# PERÍCIAS — parser do texto livre (porta do buscar-advbox-pericias.js)
# ----------------------------------------------------------------------
MESES = {"janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5,
         "junho": 6, "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10,
         "novembro": 11, "dezembro": 12}


def extrair_detalhes_pericia(texto: str) -> dict:
    """Data (dd/mm/aaaa ou 'dd de mês de aaaa'), hora (HH:MM/HHhMM) e local."""
    out = {}
    if not texto:
        return out
    m = (re.search(r"Data:\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto, re.I)
         or re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", texto))
    if m:
        out["data"] = f"{int(m.group(1)):02d}/{int(m.group(2)):02d}/{m.group(3)}"
    else:
        m2 = re.search(r"(\d{1,2})\s+de\s+([A-Za-zçãéô]+)\s+de\s+(\d{4})", texto, re.I)
        if m2:
            mes = MESES.get(pj.norm(m2.group(2)).lower())
            if mes:
                out["data"] = f"{int(m2.group(1)):02d}/{mes:02d}/{m2.group(3)}"
    mh = re.search(r"(?:às|as)\s*(\d{1,2})[:h](\d{2})", texto, re.I)
    if mh:
        out["hora"] = f"{int(mh.group(1)):02d}:{mh.group(2)}"
    ml = re.search(r"Local:\s*([^\n]+)", texto, re.I)
    if ml and ml.group(1).strip():
        out["local"] = ml.group(1).strip()
    return out


def registros_pericia(mock: bool) -> list[dict]:
    if mock:
        settings = _MOCK_SETTINGS
    elif not ADVBOX_TOKEN:
        pj.log("⚠️  ADVBOX_TOKEN vazio — use --mock ou exporte o token.")
        return []
    else:
        settings = _get_json("/settings")

    tarefas = [t for t in (settings.get("tasks") or []) if "PERICIA" in pj.norm(t.get("task"))]
    if not tarefas:
        pj.log("  (nenhum tipo de tarefa com 'perícia' no nome em /settings)")
        return []

    regs = []
    for t in tarefas:
        if mock:
            data = _MOCK_POSTS_PERICIA.get(t["id"], {})
        else:
            data = _get_json("/posts", {"task_id": t["id"], "limit": 100})
        for p in (data.get("data") or []):
            lw = p.get("lawsuit") or {}
            reg = {
                "processo": lw.get("process_number"),
                "cliente": ", ".join(c.get("name") for c in (lw.get("customers") or [])
                                     if c.get("name")),
            }
            reg.update(extrair_detalhes_pericia(p.get("notes") or ""))
            if reg.get("processo"):
                regs.append(reg)
    return regs


def registros_emenda(mock: bool) -> list[dict]:
    # «PENDENTE» — de-para real do AdvBox p/ emendas ainda não definido.
    if not mock:
        pj.log("  ⚠️  emendas: mapeamento do AdvBox ainda não definido — só --mock por ora.")
        return []
    return list(_MOCK_EMENDAS)


def registros_casos(mock: bool) -> list[dict]:
    # «PENDENTE» — de-para real do AdvBox p/ casos (status/fase) ainda não definido.
    if not mock:
        pj.log("  ⚠️  casos: mapeamento do AdvBox ainda não definido — só --mock por ora.")
        return []
    return list(_MOCK_CASOS)


# ----------------------------------------------------------------------
# infra: carrega config de cada robô, casa por processo, self-heal
# ----------------------------------------------------------------------
def _carrega_config(subpasta: str):
    caminho = os.path.join(BASE, subpasta, "config.py")
    spec = importlib.util.spec_from_file_location(f"cfg_{subpasta}", caminho)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


DESTINOS = {
    "pericia": {"sub": "avisos-pericia", "fonte": registros_pericia},
    "emenda":  {"sub": "alarme-emendas", "fonte": registros_emenda},
    "casos":   {"sub": "gatilhos-status", "fonte": registros_casos},
}


def norm_proc(v) -> str:
    return "".join(ch for ch in str(v or "") if ch.isalnum())


def merge_mae(cfg, registros: list[dict], dry: bool) -> tuple[int, int]:
    xlsx, aba, COL = cfg.MAE_XLSX, cfg.ABA, cfg.COL

    if not os.path.exists(xlsx):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = aba
        ws.append(list(COL.values()))
        wb.save(xlsx)
        pj.log(f"  (criei a mãe vazia com cabeçalho: {os.path.basename(xlsx)})")

    wb, ws = pj.abrir_mae(xlsx, aba)
    idx = pj.mapear_colunas(ws, COL)

    col_proc = idx["processo"]
    linha_de = {}
    for i in range(2, ws.max_row + 1):
        p = norm_proc(ws.cell(row=i, column=col_proc).value)
        if p:
            linha_de[p] = i

    novos = preenchidos = 0
    for reg in registros:
        p = norm_proc(reg.get("processo"))
        if not p:
            continue
        if p not in linha_de:
            i = ws.max_row + 1
            for chave, valor in reg.items():
                if chave in idx and idx[chave]:
                    ws.cell(row=i, column=idx[chave]).value = valor
            linha_de[p] = i
            novos += 1
            pj.log(f"  + novo: {reg.get('processo')} · {reg.get('cliente','')}")
        else:
            i = linha_de[p]
            faltou = []
            for chave, valor in reg.items():
                if chave not in idx or not idx[chave]:
                    continue
                atual = ws.cell(row=i, column=idx[chave]).value
                if str(atual or "").strip() == "" and str(valor or "").strip() != "":
                    ws.cell(row=i, column=idx[chave]).value = valor
                    faltou.append(COL[chave])
            if faltou:
                preenchidos += 1
                pj.log(f"  ~ completei {reg.get('processo')}: {', '.join(faltou)}")

    if not dry and (novos or preenchidos):
        pj.backup_mae(xlsx, cfg.LOG_DIR)
        wb.save(xlsx)
    return novos, preenchidos


def main() -> int:
    ap = argparse.ArgumentParser(description="PUXA-ADVBOX · preenche as mães a partir do AdvBox")
    ap.add_argument("--write", action="store_true", help="grava nas mães (padrão: --dry)")
    ap.add_argument("--dry", action="store_true")
    ap.add_argument("--mock", action="store_true", help="usa dados de exemplo, sem chamar a API")
    ap.add_argument("--destino", choices=["pericia", "emenda", "casos", "todos"], default="todos")
    args = ap.parse_args()

    dry = not args.write
    pj.log(f"PUXA-ADVBOX · modo={'DRY' if dry else 'WRITE'}"
           f"{' · MOCK' if args.mock else ''} · destino={args.destino}")

    alvos = ["pericia", "emenda", "casos"] if args.destino == "todos" else [args.destino]
    for d in alvos:
        cfg = _carrega_config(DESTINOS[d]["sub"])
        try:
            regs = DESTINOS[d]["fonte"](args.mock)
        except urllib.error.HTTPError as e:
            pj.log(f"[{d}] ❌ HTTP {e.code} no AdvBox"); continue
        except Exception as e:  # noqa: BLE001
            pj.log(f"[{d}] ❌ erro no AdvBox → {e}"); continue
        n, c = merge_mae(cfg, regs, dry)
        pj.log(f"[{d}] novos={n} · completados={c}")

    pj.log("fim" + ("  (DRY — nada gravado)" if dry else ""))
    return 0


# ----------------------------------------------------------------------
# dados de exemplo p/ --mock (no formato REAL que a API do AdvBox devolve)
# ----------------------------------------------------------------------
_MOCK_SETTINGS = {"tasks": [
    {"id": 700, "task": "PERÍCIA MÉDICA"},
    {"id": 701, "task": "AUDIÊNCIA DE INSTRUÇÃO"},
]}
_MOCK_POSTS_PERICIA = {
    700: {"data": [
        {"id": 9001, "date": "2026-07-07",
         "notes": "Perícia designada. Data: 22/07/2026 às 09:30. "
                  "Local: Av. Paulista, 1000 - Perito Dr. Silva - São Paulo/SP",
         "lawsuit": {"process_number": "0000123-45.2026.4.03.6300",
                     "customers": [{"name": "Maria Aparecida de Souza"}]}},
    ]},
}
_MOCK_EMENDAS = [
    {"processo": "0000987-65.2026.4.03.6300", "cliente": "João Carlos Ferreira",
     "intimacao": "05/07/2026", "prazo_fatal": "17/07/2026"},
]
_MOCK_CASOS = [
    {"processo": "0000123-45.2026.4.03.6300", "cliente": "Maria Aparecida de Souza",
     "status": "PERICIA DESIGNADA", "status_desde": "07/07/2026"},
    {"processo": "0000987-65.2026.4.03.6300", "cliente": "João Carlos Ferreira",
     "status": "EMENDA PENDENTE", "status_desde": "05/07/2026"},
]


if __name__ == "__main__":
    sys.exit(main())
