#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera OFICIAL_DOCUMENTOS_2026.xlsx modelo p/ testar em --dry."""
from __future__ import annotations
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill
import config as C

wb = openpyxl.Workbook(); ws = wb.active; ws.title = C.ABA
ws.append(list(C.COL.values()))
azul = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
for cell in ws[1]:
    cell.fill = azul; cell.font = Font(color="FFFFFF", bold=True)

hoje = dt.date.today()
def d(n):
    return (hoje + dt.timedelta(days=n)).strftime("%d/%m/%Y")

# proc, cliente, telefone, solicitado_em, docs_faltando, status, recebido_em, carimbos...
linhas = [
    ["001-11.2026", "Maria Souza", "5511999990001", d(-3),
     "RG ou CNH; Laudo médico atualizado; CTPS", "AGUARDANDO DOCUMENTOS", ""],
    ["002-22.2026", "João Ferreira", "5511999990002", d(-7),
     "Comprovante de residência; Exames de imagem", "AGUARDANDO DOCUMENTOS", ""],
    ["003-33.2026", "Ana Lima", "5511999990003", d(-12),
     "PPP; CAT", "AGUARDANDO DOCUMENTOS", ""],
    ["004-44.2026", "Rita Nunes", "5511999990004", d(-5),
     "Laudo", "COMPLETO", d(-1)],   # completo → silêncio
]
for base in linhas:
    ws.append(base + ["", "", "", ""])

for j, w in enumerate([16, 20, 16, 14, 42, 22, 14, 10, 10, 12, 10], start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

wb.save(C.MAE_XLSX)
print(f"✅ modelo criado: {C.MAE_XLSX}")
print("   Teste:  python3 robo_documentos.py --dry")
