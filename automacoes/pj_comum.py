# -*- coding: utf-8 -*-
"""
Jonas Inácio Automações · biblioteca comum aos robôs (perícia, emendas, gatilhos)
======================================================================

Tudo que os robôs compartilham mora aqui, para você ajustar UMA vez:
- os senders (cliente e grupo interno, TUDO via uazapi);
- datas, dias úteis e feriados;
- leitura/escrita da planilha "mãe" (openpyxl) + backup;
- estado/dedupe em JSON.

>>> Um provedor só: a MESMA instância uazapi atende cliente e grupo GERAL. Só
    falta você preencher UAZAPI_URL / UAZAPI_TOKEN / UAZAPI_GRUPO_GERAL no .env.
"""

from __future__ import annotations

import datetime as dt
import json
import os
import re
import shutil
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Rode:  pip3 install openpyxl", file=sys.stderr)
    raise


# ----------------------------------------------------------------------
# .env — carrega automaticamente automacoes/.env (fora do git), se existir.
# Assim você guarda os segredos num arquivo local em vez do ~/.zshrc.
# Não sobrescreve variável já definida no ambiente.
# ----------------------------------------------------------------------
def _carregar_dotenv():
    caminho = os.path.join(os.path.dirname(__file__), ".env")
    if not os.path.exists(caminho):
        return
    with open(caminho, encoding="utf-8") as f:
        for linha in f:
            linha = linha.strip()
            if not linha or linha.startswith("#") or "=" not in linha:
                continue
            chave, _, valor = linha.partition("=")
            chave = chave.strip()
            valor = valor.strip().strip('"').strip("'")
            os.environ.setdefault(chave, valor)


_carregar_dotenv()


# ----------------------------------------------------------------------
# log
# ----------------------------------------------------------------------
def log(msg: str) -> None:
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}")


# ----------------------------------------------------------------------
# texto / números
# ----------------------------------------------------------------------
def primeiro_nome(nome: str) -> str:
    nome = (nome or "").strip()
    return nome.split()[0].capitalize() if nome else "tudo bem"


def so_digitos(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def norm(v) -> str:
    """MAIÚSCULAS sem acento — p/ comparar STATUS de forma robusta."""
    s = unicodedata.normalize("NFKD", str(v or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.strip().upper()


# ----------------------------------------------------------------------
# datas
# ----------------------------------------------------------------------
def parse_data(v):
    """Aceita datetime do Excel, texto dd/mm/aaaa, dd/mm/aa ou aaaa-mm-dd."""
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    txt = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(txt, fmt).date()
        except ValueError:
            continue
    return None


def fmt_data(d) -> str:
    return d.strftime("%d/%m/%Y") if isinstance(d, (dt.date, dt.datetime)) else ""


# Feriados fixos p/ contagem de dias úteis. Preencha os do seu foro/ano.
# (formato "aaaa-mm-dd"). Sem isso, conta só sábado/domingo.
FERIADOS: set[dt.date] = set()


def _carregar_feriados():
    """Lê feriados de automacoes/feriados.txt (1 data por linha) se existir.
    Aceita comentário inline: '01/01/2026   # Ano Novo'."""
    caminho = os.path.join(os.path.dirname(__file__), "feriados.txt")
    if os.path.exists(caminho):
        with open(caminho, encoding="utf-8") as f:
            for ln in f:
                ln = ln.split("#", 1)[0].strip()  # tira comentário e espaços
                if not ln:
                    continue
                d = parse_data(ln)
                if d:
                    FERIADOS.add(d)


_carregar_feriados()


def eh_dia_util(d: dt.date) -> bool:
    return d.weekday() < 5 and d not in FERIADOS


def dias_uteis_entre(inicio: dt.date, fim: dt.date) -> int:
    """
    Dias ÚTEIS de `inicio` (exclusivo) até `fim` (inclusivo).
    Positivo se fim > inicio; negativo se fim < inicio (já venceu).
    """
    if inicio == fim:
        return 0
    passo = 1 if fim > inicio else -1
    d = inicio
    total = 0
    while d != fim:
        d += dt.timedelta(days=passo)
        if eh_dia_util(d):
            total += passo
    return total


# ----------------------------------------------------------------------
# estado / dedupe
# ----------------------------------------------------------------------
def carregar_estado(caminho: str) -> dict:
    if os.path.exists(caminho):
        try:
            with open(caminho, encoding="utf-8") as f:
                return json.load(f)
        except (json.JSONDecodeError, OSError):
            return {}
    return {}


def salvar_estado(caminho: str, estado: dict) -> None:
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(estado, f, ensure_ascii=False, indent=2)


# ----------------------------------------------------------------------
# planilha "mãe"
# ----------------------------------------------------------------------
def abrir_mae(xlsx: str, aba: str):
    if not os.path.exists(xlsx):
        log(f"❌ planilha mãe não encontrada: {xlsx}")
        log("   Gere a modelo com o criar_planilha_modelo.py da pasta do robô.")
        sys.exit(1)
    wb = openpyxl.load_workbook(xlsx)
    if aba not in wb.sheetnames:
        log(f"❌ aba '{aba}' não existe. Abas: {wb.sheetnames}")
        sys.exit(1)
    return wb, wb[aba]


def mapear_colunas(ws, COL: dict) -> dict:
    """Casa cabeçalho da linha 1 -> índice (1-based). Aborta se faltar coluna."""
    header = {}
    for j, cell in enumerate(ws[1], start=1):
        nome = str(cell.value or "").strip().upper()
        if nome:
            header[nome] = j
    idx, faltando = {}, []
    for chave, nome_col in COL.items():
        j = header.get(nome_col.upper())
        if j is None:
            faltando.append(nome_col)
        idx[chave] = j
    if faltando:
        log(f"❌ colunas ausentes na mãe: {', '.join(faltando)}")
        log("   Esperado: " + ", ".join(COL.values()))
        sys.exit(1)
    return idx


def backup_mae(xlsx: str, log_dir: str) -> None:
    """Copia a mãe antes de escrever (Falha grita: nunca perder a planilha)."""
    try:
        os.makedirs(log_dir, exist_ok=True)
        destino = os.path.join(log_dir, f"mae_backup_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx")
        shutil.copy2(xlsx, destino)
    except Exception as e:  # noqa: BLE001
        log(f"⚠️  não consegui fazer backup da mãe: {e}")


def carimbo_agora(hoje: dt.date) -> str:
    return f"{fmt_data(hoje)} {dt.datetime.now():%H:%M}"


# ======================================================================
# SAÍDA — senders (TUDO via uazapi: cliente e grupo interno no mesmo canal)
# ======================================================================
# Um provedor só. A MESMA instância uazapi manda tanto pro cliente (número do
# WhatsApp dele) quanto pro grupo GERAL (JID do grupo). Contrato:
#   POST {UAZAPI_URL}/send/text
#   cabeçalho:  token: <token da instância>
#   corpo JSON: {"number": "<numero ou JID>", "text": "<mensagem>"}
# Credenciais só por variável de ambiente (nunca commitar segredo).
UAZAPI_URL   = os.environ.get("UAZAPI_URL", "").rstrip("/")  # ex.: https://jonasinacioadv.uazapi.com
UAZAPI_TOKEN = os.environ.get("UAZAPI_TOKEN", "")            # token da instância (cabeçalho 'token')
UAZAPI_GRUPO = os.environ.get("UAZAPI_GRUPO_GERAL", "")      # JID do grupo, ex.: 120363...@g.us


def normalizar_telefone_br(v) -> str:
    """Só dígitos, com DDI 55. uazapi quer o número cheio (mantém o 9 do celular)."""
    d = re.sub(r"\D", "", str(v or ""))
    if not d.startswith("55") and 10 <= len(d) <= 11:
        d = "55" + d
    return d


# Seam de teste/ensaio: PJ_FAKE_SEND=1 finge um envio bem-sucedido (sem rede),
# exercitando o caminho de carimbo/dedupe. Útil p/ testes e p/ ensaiar o --send
# sem disparar mensagem de verdade. NÃO deixe ligado em produção.
_FAKE_SEND = os.environ.get("PJ_FAKE_SEND") == "1"


def _uazapi_post(number: str, text: str) -> tuple[bool, str]:
    """POST /send/text da uazapi. number = telefone do cliente OU JID do grupo."""
    if not (UAZAPI_URL and UAZAPI_TOKEN):
        return False, "uazapi sem credencial (UAZAPI_URL/UAZAPI_TOKEN)"
    import urllib.request
    import urllib.error
    body = json.dumps({"number": number, "text": text}).encode()
    req = urllib.request.Request(
        f"{UAZAPI_URL}/send/text", data=body,
        headers={"Content-Type": "application/json", "token": UAZAPI_TOKEN})
    try:
        with urllib.request.urlopen(req, timeout=30) as r:
            raw = r.read().decode("utf-8", "replace")
    except urllib.error.HTTPError as e:
        try:
            detalhe = e.read().decode("utf-8", "replace")[:160]
        except Exception:  # noqa: BLE001
            detalhe = ""
        return False, f"uazapi HTTP {e.code}: {detalhe}"
    except Exception as e:  # noqa: BLE001
        return False, f"erro uazapi: {e}"
    try:
        data = json.loads(raw)
    except ValueError:
        data = {}
    if isinstance(data, dict) and data.get("error"):
        return False, f"uazapi recusou: {data['error']}"
    return True, "uazapi OK"


def enviar_whatsapp_cliente(telefone: str, texto: str, dry: bool) -> tuple[bool, str]:
    """WhatsApp ao cliente via uazapi (/send/text para o número dele)."""
    if dry:
        return True, "DRY (não enviou)"
    if _FAKE_SEND:
        return True, "FAKE-OK"
    return _uazapi_post(normalizar_telefone_br(telefone), texto)


def enviar_alerta_interno(texto: str, dry: bool) -> tuple[bool, str]:
    """Alerta no grupo interno GERAL via uazapi (/send/text para o JID do grupo)."""
    if dry:
        return True, "DRY (não enviou)"
    if _FAKE_SEND:
        return True, "FAKE-OK"
    if not UAZAPI_GRUPO:
        return False, "uazapi sem grupo (UAZAPI_GRUPO_GERAL)"
    return _uazapi_post(UAZAPI_GRUPO, texto)
