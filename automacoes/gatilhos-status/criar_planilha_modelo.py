#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera OFICIAL_CASOS_2026.xlsx modelo p/ testar em --dry."""
from __future__ import annotations
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill
import config as C

wb = openpyxl.Workbook()
ws = wb.active
ws.title = C.ABA
ws.append(list(C.COL.values()))
azul = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
for cell in ws[1]:
    cell.fill = azul
    cell.font = Font(color="FFFFFF", bold=True)

hoje = dt.date.today()
def d(n):
    return (hoje + dt.timedelta(days=n)).strftime("%d/%m/%Y")

# proc, cliente, status, status_desde, responsavel
linhas = [
    ["001-11.2026.4.03", "Maria Souza",   "CONTRATO ASSINADO", d(0),   "Bia"],
    ["002-22.2026.4.03", "João Ferreira", "PROTOCOLO ADM",     d(-10), "Natália"],  # atrasado (SLA 3)
    ["003-33.2026.4.03", "Ana Lima",      "PASTA FECHADA",     d(-2),  "Pedro"],
    ["004-44.2026.4.03", "Pedro Alves",   "LAUDO JUNTADO",     d(-1),  "Natália"],
    ["005-55.2026.4.03", "Rita Nunes",    "EXECUCAO",          d(-3),  "Jonas"],
]
for base in linhas:
    ws.append(base)

for j, w in enumerate([20, 20, 20, 14, 14], start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

wb.save(C.MAE_XLSX)
print(f"✅ modelo criado: {C.MAE_XLSX}")
print("   Teste (1ª rodada faz linha de base):  python3 robo_status.py --dry --disparar-primeira")
