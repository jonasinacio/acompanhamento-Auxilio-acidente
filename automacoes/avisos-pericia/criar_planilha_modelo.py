#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera a planilha "mãe" modelo (OFICIAL_PERICIAS_2026.xlsx) com o cabeçalho
que o robô espera + 2 linhas de exemplo para você testar em --dry.

    python3 criar_planilha_modelo.py

Depois abra no Excel/Google Sheets, apague os exemplos e preencha com as
perícias reais. NÃO renomeie as colunas de cabeçalho.
"""
from __future__ import annotations
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill

import config as C

wb = openpyxl.Workbook()
ws = wb.active
ws.title = C.ABA

cabecalho = list(C.COL.values())
ws.append(cabecalho)

# estilo do cabeçalho (azul do manual)
azul = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
branco = Font(color="FFFFFF", bold=True)
for cell in ws[1]:
    cell.fill = azul
    cell.font = branco

# 2 exemplos: um a 15 dias, um acidentário a 7 dias (troque as datas ao testar)
hoje = dt.date.today()
ws.append([
    "0000001-11.2026.4.03.6300", "Maria Aparecida de Souza", "5511999990001",
    (hoje + dt.timedelta(days=15)).strftime("%d/%m/%Y"), "09:30",
    "Av. Paulista, 1000 - Perito Dr. Silva - São Paulo/SP",
    "Judicial", "NAO", "DESIGNADA", "", "", "", "", "",
])
ws.append([
    "0000002-22.2026.4.03.6300", "João Carlos Ferreira", "5511999990002",
    (hoje + dt.timedelta(days=7)).strftime("%d/%m/%Y"), "14:00",
    "Rua da Perícia, 50 - INSS Santo Amaro - São Paulo/SP",
    "Judicial", "SIM", "DESIGNADA", "", "", "", "", "",
])

# largura amigável
larguras = [26, 26, 16, 13, 8, 40, 12, 12, 12, 18, 18, 18, 18, 18]
for j, w in enumerate(larguras, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

wb.save(C.MAE_XLSX)
print(f"✅ planilha modelo criada: {C.MAE_XLSX}")
print(f"   aba: {C.ABA} · colunas: {', '.join(cabecalho)}")
print("   Teste:  python3 robo_pericias.py --dry")
