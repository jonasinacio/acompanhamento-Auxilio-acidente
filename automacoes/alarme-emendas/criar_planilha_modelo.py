#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera OFICIAL_EMENDAS_2026.xlsx modelo com exemplos p/ testar em --dry."""
from __future__ import annotations
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill
import config as C

wb = openpyxl.Workbook()
ws = wb.active
ws.title = C.ABA
ws.append(list(C.COL.values()))
azul = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
for cell in ws[1]:
    cell.fill = azul
    cell.font = Font(color="FFFFFF", bold=True)

hoje = dt.date.today()
def d(n):
    return (hoje + dt.timedelta(days=n)).strftime("%d/%m/%Y")

# exemplos cobrindo os níveis de alarme (ajuste as datas ao testar)
linhas = [
    # proc, cliente, intimacao, prazo_fatal, classif, dep_doc, doc_ok, status, protoc_em, carimbos...
    ["001-11.2026.4.03", "Maria Souza",   d(-1), d(10), "",           "SIM", "NAO", "EMENDA PENDENTE", ""],
    ["002-22.2026.4.03", "João Ferreira", d(-3), d(4),  "documental", "SIM", "SIM", "EMENDA PENDENTE", ""],
    ["003-33.2026.4.03", "Ana Lima",      d(-5), d(1),  "jurídica",   "NAO", "NAO", "EMENDA PENDENTE", ""],
    ["004-44.2026.4.03", "Pedro Alves",   d(-6), d(-2), "simples",    "NAO", "NAO", "EMENDA PENDENTE", ""],
    ["005-55.2026.4.03", "Rita Nunes",    d(-4), d(3),  "documental", "SIM", "SIM", "PROTOCOLADA",     d(0)],
]
for base in linhas:
    ws.append(base + ["", "", "", "", ""])  # 5 colunas de carimbo vazias

larguras = [20, 20, 12, 12, 14, 11, 8, 18, 14, 14, 14, 10, 10, 10]
for j, w in enumerate(larguras, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w

wb.save(C.MAE_XLSX)
print(f"✅ modelo criado: {C.MAE_XLSX}")
print("   Teste:  python3 robo_emendas.py --dry")
